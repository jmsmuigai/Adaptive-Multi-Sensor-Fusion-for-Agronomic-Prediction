"""
build_docx_final.py -- Render manuscript_final.md into the JAGST submission
document, placing every figure and table INLINE at its [[FIGURE n]] /
[[TABLE n]] marker rather than collecting them at the end.

JAGST requirements implemented: Calibri (body) 11 pt, single spacing,
numbered sections, illustrations placed within the text at the appropriate
points, Harvard citations with a resolvable link beneath every reference.
"""

import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
EA = os.path.join(HERE, 'ea')
FIGS = os.path.join(HERE, 'figures_ea')
OUT = os.path.join(HERE, '..', 'Mburu_et_al_JAGST_RESUBMISSION_FINAL.docx')

# manuscript figure number -> generated file  (order follows first reference)
FIG_FILE = {
    1: 'Figure_1_study_area.png',
    2: 'Figure_2_architecture.png',
    3: 'Figure_3_cluster_cv.png',
    4: 'Figure_4_obs_pred.png',
    5: 'Figure_5_designs.png',
    6: 'Figure_6_fold_spread.png',
    7: 'Figure_7_country.png',
    8: 'Figure_8_residuals.png',
    9: 'Figure_9_attention.png',
    10: 'Figure_10_spectral_context.png',
    11: 'Figure_11_agronomy.png',
}
FIG_CAP = {
    1: 'Fig. 1. Study area, sampling structure and mid-infrared spectra. (a) The 16 '
       'AfSIS Land Degradation Surveillance Framework sentinel sites contributing to '
       'the analysis set, coloured by country. (b) Samples and sentinel sites per '
       'country. (c) Mean scatter-corrected MIR spectra by country.',
    2: 'Fig. 2. Proposed Hybrid Attention-Weighted CNN–DAFM–RBF architecture. The '
       'spectral branch applies two 1D convolutional layers with ReLU activation and '
       'max-pooling along the wavenumber axis; the Dynamic Attention Fusion Module '
       'computes additive attention over spectral positions to form a context vector; '
       'a learned gate balances the spectral and ancillary branches; and a '
       'differentiable radial basis function head produces the prediction. All '
       'operations are differentiable, so the network trains end to end.',
    3: 'Fig. 3. Within-landscape prediction under leave-cluster-out cross-validation. '
       '(a) R² by model and soil property. (b) Mean R² across the thirteen properties; '
       'the proposed architecture is shown in red. Error bars are the mean fold '
       'standard deviation.',
    4: 'Fig. 4. Out-of-fold observed versus predicted values for the four '
       'best-predicted properties, pooled across leave-cluster-out evaluation folds. '
       'The red dashed line is the 1:1 line of perfect agreement.',
    5: 'Fig. 5. Predictive skill under three nested spatial validation designs. '
       '(a) Mean R² for each model as the withheld unit widens from cluster to site to '
       'country. (b) Leave-site-out R² by property. (c) Number of properties retaining '
       'positive R² under each design.',
    6: 'Fig. 6. Fold-level spread of R² under leave-cluster-out cross-validation. '
       'Vertical bars span the four evaluation folds and markers give fold means, so '
       'bar width indicates how stable each estimate is.',
    7: 'Fig. 7. Leave-country-out transfer resolved by held-out country. (a) Mean R² '
       'for each model when each country is withheld. (b) Best R² per property and '
       'country. (c) Relationship between mean skill and the number of properties '
       'retaining positive R². Ethiopia served as the tuning fold and is not reported '
       'as a test country.',
    8: 'Fig. 8. Residual diagnostics for the three best-predicted properties. Top row: '
       'residuals against fitted values, testing for systematic bias across the '
       'prediction range. Bottom row: residual distributions separated by country, '
       'testing whether error concentrates in one national dataset.',
    9: 'Fig. 9. Interpretability of the Dynamic Attention Fusion Module. (a) Mean '
       'attention as a function of wavenumber, with the uniform-attention reference; '
       'shaded bands are ±1 standard deviation across folds. (b) Learned gate weights '
       'showing the measured contribution of the spectral and ancillary branches.',
    10: 'Fig. 10. Attention placed in spectroscopic context. (a) Mean scatter-corrected '
       'spectrum with principal absorbance regions shaded. (b) Mean DAFM attention on '
       'the same wavenumber axis, showing where learned weights fall relative to known '
       'functional-group absorbance.',
    11: 'Fig. 11. Agronomic status and predictability. (a) Percentage of topsoils below '
        'indicative critical levels, with whiskers at ±50% of each threshold. '
        '(b) Deficiency prevalence against best achievable R².',
}
# manuscript table number -> key in ea/tables.json
TBL_KEY = {1: 'table1', 2: 'table2', 3: 'table3', 4: 'table4', 5: 'table5',
           6: 'table_bound', 7: 'table6', 8: 'table10', 9: 'table9',
           10: 'table_timing', 11: 'table8', 12: 'table7'}

MARK = re.compile(r'^\[\[(FIGURE|TABLE) (\d+)\]\]$')
RICH = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*')


def set_style(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), 'Calibri')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.0
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)


def page_numbers(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for el, attr in (('w:fldChar', {'w:fldCharType': 'begin'}),
                         ('w:instrText', None),
                         ('w:fldChar', {'w:fldCharType': 'end'})):
            e = OxmlElement(el)
            if attr:
                for k, v in attr.items():
                    e.set(qn(k), v)
            else:
                e.set(qn('xml:space'), 'preserve')
                e.text = 'PAGE'
            r._r.append(e)
        r.font.size = Pt(9)


def rich(par, text, size=11):
    pos = 0
    for m in RICH.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            r.font.size = Pt(size)
        r = par.add_run(m.group(1) or m.group(2))
        r.font.size = Pt(size)
        r.bold = bool(m.group(1))
        r.italic = bool(m.group(2))
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        r.font.size = Pt(size)
    return par


def put_table(doc, spec, size=7.5):
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(10)
    cap.paragraph_format.keep_with_next = True
    rich(cap, spec['caption'], size=9)
    rows = spec['rows']
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fs = size if ncol <= 9 else 6.5
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(fs)
            r.font.name = 'Calibri'
            if i == 0 or str(row[0]) in ('Mean', 'Total', 'Mean R²'):
                r.bold = True
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def put_figure(doc, num):
    path = os.path.join(FIGS, FIG_FILE[num])
    if not os.path.exists(path):
        print('  MISSING FIGURE', num, FIG_FILE[num])
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(6.3))
    cap = doc.add_paragraph()
    r = cap.add_run(FIG_CAP[num])
    r.font.size = Pt(9)
    doc.add_paragraph()


def main():
    md = open(os.path.join(HERE, 'manuscript_final.md')).read().split('\n')
    tables = json.load(open(os.path.join(EA, 'tables.json')))

    doc = Document()
    set_style(doc)
    page_numbers(doc)
    in_refs = False

    for line in md:
        line = line.rstrip()
        m = MARK.match(line.strip())
        if m:
            kind, num = m.group(1), int(m.group(2))
            if kind == 'FIGURE':
                put_figure(doc, num)
            else:
                put_table(doc, tables[TBL_KEY[num]])
            continue
        if not line.strip() or line.strip() == '---':
            continue
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            r = p.add_run(line[5:])
            r.bold = True
            r.italic = True
            r.font.size = Pt(10.5)
            continue
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            r = p.add_run(line[4:])
            r.bold = True
            r.font.size = Pt(11)
            continue
        if line.startswith('## '):
            title = line[3:]
            in_refs = title.startswith('8.0')
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(12.5)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:])
            r.bold = True
            r.font.size = Pt(15)
            continue

        p = doc.add_paragraph()
        if in_refs and line.strip().startswith('http'):
            r = p.add_run(line.strip())
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(8)
        elif in_refs:
            rich(p, line)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(0)
        else:
            if line.lstrip().startswith(('- ', '1. ', '2. ', '3. ', '4. ', '5. ')):
                p.paragraph_format.left_indent = Inches(0.3)
            rich(p, line.lstrip())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(OUT)
    print('wrote', os.path.abspath(OUT))


if __name__ == '__main__':
    main()
