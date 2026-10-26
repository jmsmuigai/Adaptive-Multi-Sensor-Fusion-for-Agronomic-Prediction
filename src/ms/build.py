"""build.py -- render the revised manuscript (and supplementary file) to .docx.

1. merge the Markdown parts, fill [[number]] placeholders from out/numbers.json,
   insert tables and figures
2. check every in-text citation has a reference and vice versa
3. pandoc -> docx with the JAGST reference styles
4. post-process: table formatting, figure sizing, line numbers, page numbers,
   blank author metadata (blind review)
"""
import json
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, 'out')
FIG = os.path.join(OUT, os.environ.get('FIGDIR', 'fig'))
MT = os.path.join(OUT, 'ms_tables')
DOCS = os.path.join(OUT, 'docs'); os.makedirs(DOCS, exist_ok=True)
N = json.load(open(os.path.join(OUT, 'numbers.json')))
FIGS = {'FIG1': 'Figure_1_study_area.png', 'FIG2': 'Figure_2_architecture.png', 'FIG3': 'Figure_3_designs.png',
        'FIG4': 'Figure_4_site_properties.png', 'FIG5': 'Figure_5_obs_pred.png', 'FIG6': 'Figure_6_ablation.png',
        'FIG7': 'Figure_7_size_diversity.png', 'FIG8': 'Figure_8_spiking.png', 'FIG9': 'Figure_9_conformal.png',
        'FIG10': 'Figure_10_interpretability.png'}


def fill(text):
    def rep(m):
        key, fmt = m.group(1), m.group(2)
        if key == 'TABLE':
            return open(os.path.join(MT, fmt + '.md')).read()
        if key == 'REFERENCES':
            return open(os.path.join(HERE, 'references.md')).read()
        if key not in N:
            raise KeyError(f'missing number: {key}')
        v = N[key]
        if fmt:
            s = format(v, fmt)
        else:
            s = str(v)
        return s.replace('-', '−') if isinstance(v, (int, float)) else s
    text = re.sub(r'\[\[([A-Za-z0-9_:.+%,]+?)(?::([^\]]+))?\]\]', lambda m: rep(m) if not m.group(1).startswith('TABLE') else rep(m), text)
    for k, f in FIGS.items():
        text = text.replace(f'![]({k})', f'![]({os.path.join(FIG, f)}){{width=16cm}}')
    return text


def check_citations(text):
    body, refs = text.split('**8.0 References**')
    cited = set()
    for m in re.finditer(r'([A-ZÅÄÖÉØ][\w\-’\'ąćęłńóśźżÅåÄäÖöéüŽž]+(?: and [A-Z][\w\-]+)?(?: \*et al\.\*)?)(?:,| \(|\s\()?\s*\(?(\d{4})', body):
        pass
    # simpler: extract (Name, year) and "Name (year)" patterns
    pats = re.findall(r"((?:van )?[A-ZÅÄÖÉ][A-Za-zÀ-ž’'\-]+)(?: and (?:van )?[A-ZÅÄÖÉ][A-Za-zÀ-ž\-]+)?(?: \*et al\.\*)?,? \(?(\d{4})[a-z]?", body)
    refnames = re.findall(r"^((?:van )?[A-ZÅÄÖÉ][A-Za-zÀ-ž’'\- ]+?), [A-Z]\..*?\((\d{4})\)", refs, flags=re.M)
    refset = {(a.split()[-1] if not a.startswith('van ') else 'van Amersfoort', y) for a, y in refnames}
    refset |= {(a.split(',')[0].strip(), y) for a, y in refnames}
    missing = sorted({(a, y) for a, y in pats if (a, y) not in refset and a not in
                      ('Table', 'Figure', 'Section', 'Fig', 'Phase', 'Proposition', 'Kenya', 'Uganda', 'Tanzania', 'Ethiopia', 'In', 'January')})
    return missing, refset


def set_cell_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge, val in [('top', 'single'), ('left', 'nil'), ('bottom', 'single'), ('right', 'nil'),
                      ('insideH', 'nil'), ('insideV', 'nil')]:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), val)
        if val == 'single':
            el.set(qn('w:sz'), '8'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    tblPr.append(borders)


def postprocess(path, line_numbers=True, landscape=False):
    d = Document(path)
    TW = 13600 if landscape else 9070
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec0 = d.sections[0]
        sec0.orientation = WD_ORIENT.LANDSCAPE
        sec0.page_width, sec0.page_height = Cm(29.7), Cm(21.0)
    cp = d.core_properties
    cp.author = ''; cp.last_modified_by = ''; cp.title = ''; cp.comments = ''
    for t in d.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        ncol = len(t.columns)
        body = [0] * ncol; head = [0] * ncol
        for i, row in enumerate(t.rows):
            for j, c in enumerate(row.cells[:ncol]):
                txt = c.text.strip()
                if i == 0:
                    head[j] = max((len(w) for w in txt.split()), default=1)
                else:
                    body[j] = max(body[j], min(len(txt), 40))
        wts = [max(b, min(h, 11), 5) + 4 for b, h in zip(body, head)]
        widths = [int(TW * w / sum(wts)) for w in wts]
        for old in tblPr.findall(qn('w:tblW')):
            tblPr.remove(old)
        tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(TW)); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
        for old in tblPr.findall(qn('w:tblLayout')):
            tblPr.remove(old)
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        for old in tblPr.findall(qn('w:tblCellMar')):
            tblPr.remove(old)
        mar = OxmlElement('w:tblCellMar')
        for side, w in (('top', '0'), ('left', '40'), ('bottom', '0'), ('right', '40')):
            e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), w); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblPr.append(mar)
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(w))
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                tcPr = c._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(old)
                cw = OxmlElement('w:tcW'); cw.set(qn('w:w'), str(w)); cw.set(qn('w:type'), 'dxa'); tcPr.insert(0, cw)
        set_cell_borders(t)
        nrows = len(t.rows)
        for i, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit'); trPr.append(cs)
            if i == 0:
                th = OxmlElement('w:tblHeader'); trPr.append(th)
            if nrows <= 30 and i < nrows - 1:
                for c in row.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
        for i, row in enumerate(t.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        if i == 0:
                            r.font.bold = True
            if i == 0:
                for c in row.cells:
                    tcPr = c._tc.get_or_add_tcPr()
                    b = OxmlElement('w:tcBorders'); bt = OxmlElement('w:bottom')
                    bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '6'); bt.set(qn('w:color'), '000000')
                    b.append(bt); tcPr.append(b)
    for p in d.paragraphs:
        if p.runs and any(r._element.findall('.//' + qn('w:drawing')) for r in p.runs):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
        txt = p.text.strip()
        if txt.startswith('EQ::'):
            for r in p.runs:
                if 'EQ::' in r.text:
                    r.text = r.text.replace('EQ:: ', '').replace('EQ::', '')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            txt = p.text.strip()
        if re.match(r'^(Table \d+|Fig\. \d+|Table S\d+)[:.]', txt):
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            for r in p.runs[:1]:
                pass
            if txt.startswith('Table'):
                p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.size = Pt(10)
            # bold the label
            lab = re.match(r'^(Table S?\d+:|Fig\. \d+\.)', txt).group(1)
            if p.runs and p.runs[0].text.startswith(lab):
                r0 = p.runs[0]
                rest = r0.text[len(lab):]
                r0.text = lab; r0.font.bold = True
                nr = p.add_run(rest); nr.font.size = Pt(10)
                p._p.remove(nr._r); r0._r.addnext(nr._r)
        if txt.startswith('Mid-infrared (MIR) spectroscopy can lower'):
            pass
    ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
             'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
             'tblCaption', 'tblDescription']
    for t in d.tables:
        tblPr = t._tbl.tblPr
        kids = list(tblPr)
        kids.sort(key=lambda e: ORDER.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in ORDER else 99)
        for k in kids:
            tblPr.remove(k)
        for k in kids:
            tblPr.append(k)
        TC = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
              'textDirection', 'tcFitText', 'vAlign', 'hideMark']
        for row in t.rows:
            for c in row.cells:
                tcPr = c._tc.get_or_add_tcPr()
                kids = sorted(list(tcPr), key=lambda e: TC.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in TC else 99)
                for k in kids:
                    tcPr.remove(k)
                for k in kids:
                    tcPr.append(k)
    sec = d.sections[0]
    sec.header_distance = Cm(1.25); sec.footer_distance = Cm(1.25); sec.gutter = Cm(0)
    if line_numbers:
        ln = OxmlElement('w:lnNumType')
        ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous'); ln.set(qn('w:distance'), '284')
        sec._sectPr.append(ln)
    # page number footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sup = OxmlElement('w:suppressLineNumbers'); ppr = fp._p.get_or_add_pPr()
    ps = ppr.find(qn('w:pStyle'))
    if ps is not None:
        ps.addnext(sup)
    else:
        ppr.insert(0, sup)
    r = fp.add_run()
    for tag, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if tag:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag); r._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)
    d.save(path)


def finalise(path, running_title, refs_heading='8.0 References'):
    """Submission copy (JAGST checklist): Calibri 11 pt, single spacing, 6 pt after paragraphs,
    justified body text, hanging-indent references, running title in the page header."""
    d = Document(path)
    in_refs = False
    for p in d.paragraphs:
        txt = p.text.strip()
        if refs_heading and txt.startswith(refs_heading):
            in_refs = True
            continue
        if not txt or re.match(r'^(Table S?\d+:|Fig\. \d+\.)', txt):
            continue
        if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            continue
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_after = Pt(6)
        if in_refs:
            pf.left_indent = Cm(0.75); pf.first_line_indent = Cm(-0.75)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif not (p.runs and all(r.bold for r in p.runs if r.text.strip())):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if p.runs and all(r.bold for r in p.runs if r.text.strip()) and len(txt) < 120:
            pf.keep_with_next = True
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
    sec = d.sections[0]
    if running_title:
        hp = sec.header.paragraphs[0]
        hp.text = ''
        r = hp.add_run(running_title); r.italic = True; r.font.size = Pt(9)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.save(path)


def render(md_parts, out_docx, line_numbers=True, landscape=False):
    text = '\n\n'.join(open(os.path.join(HERE, p)).read() for p in md_parts)
    text = fill(text)
    md = out_docx.replace('.docx', '.md')
    open(md, 'w').write(text)
    subprocess.run(['pandoc', md, '-f', 'markdown+tex_math_dollars+superscript+subscript+pipe_tables-auto_identifiers',
                    '--reference-doc', os.path.join(HERE, 'reference.docx'), '-o', out_docx], check=True)
    postprocess(out_docx, line_numbers, landscape)
    return text


if __name__ == '__main__':
    text = render(['part0_front.md', 'part1_intro_methods.md', 'part2_results.md', 'part3_end.md'],
                  os.path.join(DOCS, 'manuscript_clean.docx'))
    missing, refset = check_citations(text)
    print('citations without reference:', missing)
    body = text.split('**8.0 References**')[0]
    unused = sorted({a for a, y in refset if a not in body and len(a) > 2})
    print('reference names not found in body:', unused)
    words = len(re.findall(r'\w+', body.split('**1.0 Introduction**')[1]))
    abst = len(re.findall(r'\w+', text.split('**Abstract**')[1].split('**Keywords')[0]))
    print('main-text words', words, 'abstract words', abst, f'({100 * abst / words:.1f}% of text)')
